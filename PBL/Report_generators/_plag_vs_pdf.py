import re
import sys

from docx import Document
from pypdf import PdfReader

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

def norm_words(text):
    return re.findall(r"[a-z0-9']+", text.lower())

# --- digital__.pdf text
pdf = PdfReader(r"C:\aman\dataset for pbl\digital__.pdf")
pdf_words = []
for page in pdf.pages:
    pdf_words.extend(norm_words(page.extract_text() or ""))
print("pdf words:", len(pdf_words))

def ngrams(ws, n):
    return set(tuple(ws[i:i + n]) for i in range(len(ws) - n + 1))

pdf6 = ngrams(pdf_words, 6)
pdf8 = ngrams(pdf_words, 8)

d = Document(r"C:\aman\dataset for pbl\A.docx")

# check each paragraph of A.docx against the pdf; report matching paragraphs
flagged = []
for i, p in enumerate(d.paragraphs):
    ws = norm_words(p.text)
    if len(ws) < 8:
        continue
    g = ngrams(ws, 6)
    hit = g & pdf6
    if hit:
        ratio = len(hit) / len(g)
        flagged.append((i, ratio, len(hit), p.text.strip()[:90]))

print("\nparagraphs sharing 6-word phrases with digital__.pdf:", len(flagged))
for i, r, n, t in sorted(flagged, key=lambda x: -x[1])[:40]:
    print(f"  para {i}  {r:.0%} ({n} phrases)  {t!r}")

# overall
all_words = []
for p in d.paragraphs:
    all_words.extend(norm_words(p.text))
a8 = ngrams(all_words, 8)
print("\noverall 8-gram overlap A.docx vs digital__.pdf:", f"{len(a8 & pdf8)/max(1,len(a8)):.2%}", f"({len(a8 & pdf8)} phrases)")
